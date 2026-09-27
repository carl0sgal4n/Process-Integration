"""PINCH LAB: integración energética con Cp constante.

Unidades: temperaturas en °C; CP = m·cp en kJ/(h·K); Q en kJ/h.
1 kW = 3600 kJ/h. No confundir CP con el calor específico cp.

Ejecutar: python pinch.py --csv corrientes_ejemplo.csv --dtmin 10
Interfaz local: python app.py

Método:
  1. Temperaturas desplazadas: T*hot=T-dTmin/2; T*cold=T+dTmin/2.
  2. Cascada por intervalos: dH=(sum CP_hot-sum CP_cold)*dT*.
  3. QHmin=max(0,-min(cascada)); QCmin=cascada_final+QHmin.
  4. Curvas compuestas: integral de CP activo frente a T.
  5. Desplazar H de la curva fría en QCmin y subdividir el eje H.
  6. En cada tramo, repartir Q entre todas las parejas activas mediante
     fracciones de CP. Cada pareja constituye un intercambiador de ramas
     a contracorriente; se comprueban sus dos diferencias de temperatura.

Se obtiene una red realizable ideal mediante división y mezcla isotérmica
de ramas. NO se optimiza el número de equipos, las pérdidas de presión ni
el coste. No se modela calor latente: una corriente isotérmica es rechazada.
"""
from __future__ import annotations

import argparse
import io
import json
import math
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


@dataclass(frozen=True)
class Stream:
    """Una corriente sensible. El tipo se deduce de Tin y Tout."""
    id: str
    cp: float
    tin: float
    tout: float

    @property
    def hot(self):
        return self.tin > self.tout

    @property
    def low(self):
        return min(self.tin, self.tout)

    @property
    def high(self):
        return max(self.tin, self.tout)

    @property
    def duty(self):
        return self.cp * (self.high - self.low)


EXAMPLE = [Stream("H1", 1000, 250, 120), Stream("H2", 4000, 200, 100),
           Stream("C3", 3000, 90, 150), Stream("C4", 6000, 130, 190)]
HOT, COLD, GREEN, NAVY = "#D64D44", "#2275C9", "#148A83", "#142D4E"


def validate(streams, dtmin, u=0):
    """Rechaza datos ambiguos antes de calcular, sin inventar ceros."""
    if not streams:
        raise ValueError("Introduce al menos una corriente.")
    if not math.isfinite(dtmin) or dtmin < 0:
        raise ValueError("ΔT mínimo debe ser finito y mayor o igual que cero.")
    if not math.isfinite(u) or u < 0:
        raise ValueError("U debe ser finito y no negativo. Usa 0 para no estimar áreas.")
    ids = [s.id.strip() for s in streams]
    if any(not x for x in ids) or len(set(ids)) != len(ids):
        raise ValueError("Los identificadores deben ser únicos y no estar vacíos.")
    for s in streams:
        if not all(math.isfinite(v) for v in (s.cp, s.tin, s.tout)):
            raise ValueError(f"{s.id}: hay números vacíos o no finitos.")
        if s.cp <= 0:
            raise ValueError(f"{s.id}: CP debe ser positivo.")
        if s.low < -273.15:
            raise ValueError(f"{s.id}: temperatura inferior al cero absoluto.")
        if abs(s.tin - s.tout) < 1e-10:
            raise ValueError(f"{s.id}: Tin=Tout. El modelo no incluye calor latente.")


def read_streams(path):
    """Lee CSV o la hoja Entradas de un Excel generado por este programa.

    Para Excel se leen SOLO valores de entrada, nunca resultados almacenados.
    El fichero debe mantener la tabla en A11:D... y los parámetros en B4/B5.
    """
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Entradas"]
        rows = []
        for row in ws.iter_rows(min_row=11, max_col=4, values_only=True):
            if row[0] is None:
                break
            rows.append(Stream(str(row[0]), *map(float, row[1:])))
        dt, u = float(ws["B4"].value), float(ws["B5"].value)
        wb.close()
        return rows, dt, u
    frame = pd.read_csv(path)
    required = ["id", "cp", "tin", "tout"]
    if not set(required).issubset(frame.columns):
        raise ValueError("El CSV necesita columnas: id,cp,tin,tout (punto decimal).")
    return [Stream(str(r.id), float(r.cp), float(r.tin), float(r.tout))
            for r in frame[required].itertuples(index=False)], None, None


def targets(streams, dtmin):
    """Problem Table Algorithm. Conserva temperaturas repetidas para Excel.

    Los intervalos de anchura cero no aportan calor y no alteran la cascada.
    Se devuelven todos los nodos de cascada cero. Los extremos pueden ser
    pinches de umbral, y un intervalo cero puede representar una zona pinch.
    """
    validate(streams, dtmin)
    shifted = [(s.low + (-dtmin/2 if s.hot else dtmin/2),
                s.high + (-dtmin/2 if s.hot else dtmin/2)) for s in streams]
    levels = np.sort(np.array(shifted).ravel())[::-1]
    mids = (levels[:-1] + levels[1:])/2
    hot_cp, cold_cp = [], []
    for t in mids:
        hot_cp.append(sum(s.cp for s, (lo, hi) in zip(streams, shifted)
                          if s.hot and lo <= t <= hi))
        cold_cp.append(sum(s.cp for s, (lo, hi) in zip(streams, shifted)
                           if not s.hot and lo <= t <= hi))
    dh = (np.array(hot_cp)-np.array(cold_cp)) * (levels[:-1]-levels[1:])
    raw = np.r_[0.0, np.cumsum(dh)]
    qh = max(0.0, -float(raw.min()))
    residual = raw + qh
    qc = float(residual[-1])
    qhot = sum(s.duty for s in streams if s.hot)
    qcold = sum(s.duty for s in streams if not s.hot)
    tol = max(1.0, qhot, qcold)*1e-9
    zero_levels = sorted(set(float(t) for t, r in zip(levels, residual)
                             if abs(r) <= tol), reverse=True)
    pinches = [{"T_star": t, "T_hot": t+dtmin/2, "T_cold": t-dtmin/2,
                "kind": "interior" if levels[-1] < t < levels[0] else "umbral"}
               for t in zero_levels]
    return dict(levels=levels, hot_cp=np.array(hot_cp), cold_cp=np.array(cold_cp),
                dh=dh, raw=raw, residual=residual, qh=qh, qc=qc,
                qhot=qhot, qcold=qcold, recovered=qhot-qc, pinches=pinches,
                tolerance=tol)


def composite(streams, hot, temperatures, offset=0):
    """Integral CP dT. Los huecos sin corrientes tienen dH=0."""
    ts = np.array(temperatures, dtype=float)
    cps = np.array([sum(s.cp for s in streams if s.hot == hot and
                        s.low <= (lo+hi)/2 <= s.high)
                    for lo, hi in zip(ts[:-1], ts[1:])])
    hs = np.r_[0., np.cumsum(cps*np.diff(ts))] + offset
    return dict(t=ts, cp=cps, h=hs)


def _point_on_segment(curve, energy_mid, energy_lo, energy_hi):
    """Localiza el tramo positivo por su punto medio: evita interpolar huecos."""
    h = curve["h"]
    indices = np.flatnonzero((h[:-1] <= energy_mid) & (energy_mid < h[1:]))
    if len(indices) == 0:
        return None
    k = int(indices[0])
    cp = float(curve["cp"][k])
    low = float(curve["t"][k] + (energy_lo-h[k])/cp)
    high = float(curve["t"][k] + (energy_hi-h[k])/cp)
    return dict(index=k+1, low=low, high=high, cp=cp)


def lmtd(a, b):
    """Media logarítmica estable. Un extremo cero exige área infinita."""
    if min(a, b) <= 1e-10:
        return 0.0
    if math.isclose(a, b, rel_tol=1e-9):
        return (a+b)/2
    return (a-b)/math.log(a/b)


def analyse(streams, dtmin=10.0, u=0.0):
    """Calcula objetivos, curvas y una red ideal que alcanza QHmin/QCmin.

    u: coeficiente global OPCIONAL en kJ/(h·m²·K); 0 significa desconocido.
    Para un tramo de energía q, las fracciones de asignación son
      fh_i = CP_i / sum(CP_hot), fc_j = CP_j / sum(CP_cold).
    Se asigna Q_ij=q*fh_i*fc_j.
    La rama de la corriente caliente i lleva fracción fc_j de su caudal;
    la rama fría j lleva fracción fh_i. No confundir estas dos fracciones.
    Se mezclan ramas de una misma corriente a igual T al terminar cada tramo.
    """
    validate(streams, dtmin, u)
    r = targets(streams, dtmin)
    ts = sorted(v for s in streams for v in (s.low, s.high))
    hot = composite(streams, True, ts)
    cold = composite(streams, False, ts, r["qc"])
    # Nodos de las dos curvas sobre un mismo eje de entalpía.
    energies = np.sort(np.r_[hot["h"], cold["h"]])
    blocks, exchangers, utilities = [], [], []
    matrix = np.zeros((len(streams), len(streams)))
    heating, cooling = np.zeros(len(streams)), np.zeros(len(streams))
    for k, (elo, ehi) in enumerate(zip(energies[:-1], energies[1:])):
        q = float(ehi-elo)
        hm = _point_on_segment(hot, (elo+ehi)/2, elo, ehi) if q > r["tolerance"] else None
        cm = _point_on_segment(cold, (elo+ehi)/2, elo, ehi) if q > r["tolerance"] else None
        fh = np.zeros(len(streams)); fc = np.zeros(len(streams))
        for i, s in enumerate(streams):
            if hm and s.hot and s.low <= (hm["low"]+hm["high"])/2 <= s.high:
                fh[i] = s.cp/hm["cp"]
            if cm and not s.hot and s.low <= (cm["low"]+cm["high"])/2 <= s.high:
                fc[i] = s.cp/cm["cp"]
        block = dict(index=k+1, elo=float(elo), ehi=float(ehi), q=q,
                     hot=hm, cold=cm, fh=fh, fc=fc,
                     process=q if hm and cm else 0.,
                     cooling=q if hm and not cm else 0.,
                     heating=q if cm and not hm else 0.)
        blocks.append(block)
        if hm and cm:
            d1, d2 = hm["high"]-cm["high"], hm["low"]-cm["low"]
            if min(d1, d2) < dtmin-1e-6:
                raise ArithmeticError("La asignación viola ΔT mínimo.")
            for i in np.flatnonzero(fh):
                for j in np.flatnonzero(fc):
                    duty = q*fh[i]*fc[j]
                    matrix[i, j] += duty
                    logmean = lmtd(d1, d2)
                    exchangers.append(dict(equipo=f"E{len(exchangers)+1:02d}",
                        tramo=k+1, caliente=streams[i].id, fria=streams[j].id,
                        Q_kJh=duty, Th_in=hm["high"], Th_out=hm["low"],
                        Tc_in=cm["low"], Tc_out=cm["high"],
                        fraccion_rama_H=float(fc[j]), fraccion_rama_C=float(fh[i]),
                        CP_rama_H=streams[i].cp*fc[j], CP_rama_C=streams[j].cp*fh[i],
                        DT_extremo_1=d1, DT_extremo_2=d2, DT_lm=logmean,
                        area_m2=(duty/(u*logmean) if logmean > 0 else math.inf) if u else None))
        elif hm or cm:
            fractions = fh if hm else fc
            for i in np.flatnonzero(fractions):
                duty = q*fractions[i]
                if hm: cooling[i] += duty
                else: heating[i] += duty
                p = hm if hm else cm
                utilities.append(dict(tramo=k+1, corriente=streams[i].id,
                    servicio="Refrigeración" if hm else "Calentamiento", Q_kJh=duty,
                    T_in=p["high"] if hm else p["low"],
                    T_out=p["low"] if hm else p["high"]))
    # Balances independientes por corriente: toda la demanda debe estar asignada.
    errors = []
    for i, s in enumerate(streams):
        assigned = matrix[i, :].sum()+cooling[i] if s.hot else matrix[:, i].sum()+heating[i]
        errors.append(float(assigned-s.duty))
    if max(abs(v) for v in errors) > r["tolerance"]*10:
        raise ArithmeticError("No cierra el balance de alguna corriente.")
    if abs(heating.sum()-r["qh"]) > r["tolerance"]*10 or abs(cooling.sum()-r["qc"]) > r["tolerance"]*10:
        raise ArithmeticError("La red no coincide con los objetivos de servicios.")
    r.update(streams=streams, dtmin=float(dtmin), u=float(u), hot=hot, cold=cold,
             energies=energies, blocks=blocks, exchangers=exchangers,
             utilities=utilities, matrix=matrix, heating=heating, cooling=cooling,
             balance_errors=errors)
    return r


def tables(r):
    """Tablas en pandas para Jupyter/Colab, CSV y la interfaz."""
    s = r["streams"]
    return {
        "corrientes": pd.DataFrame([dict(id=x.id, tipo="HOT" if x.hot else "COLD",
            CP_kJhK=x.cp, Tin_C=x.tin, Tout_C=x.tout, Q_kJh=x.duty,
            Q_kW=x.duty/3600) for x in s]),
        "cascada": pd.DataFrame(dict(T_superior=r["levels"][:-1],
            T_inferior=r["levels"][1:], CP_hot=r["hot_cp"], CP_cold=r["cold_cp"],
            delta_H=r["dh"], residual_sin_ajustar=r["raw"][1:],
            residual_ajustado=r["residual"][1:])),
        "pinches": pd.DataFrame(r["pinches"]),
        "intercambiadores": pd.DataFrame(r["exchangers"]),
        "servicios": pd.DataFrame(r["utilities"]),
        "matriz_Q": pd.DataFrame(r["matrix"], index=[x.id for x in s], columns=[x.id for x in s]),
        "balances": pd.DataFrame(dict(corriente=[x.id for x in s],
            refrigeracion=r["cooling"], calentamiento=r["heating"], error=r["balance_errors"]))}


def sensitivity(streams, dt_values):
    """Recalcula cada escenario de ΔT mínimo; no desplaza una curva a ojo."""
    rows = []
    for dt in dt_values:
        r = targets(streams, float(dt))
        rows.append(dict(DTmin=float(dt), QHmin=r["qh"], QCmin=r["qc"], Qrec=r["recovered"]))
    return pd.DataFrame(rows)


def figures(r):
    """Gráficos científicos independientes, exportables a PNG, SVG o PDF."""
    plt.rcParams.update({"font.family":"DejaVu Sans", "font.size":10,
        "axes.spines.top":False, "axes.spines.right":False,
        "axes.titleweight":"bold", "axes.titlecolor":NAVY,
        "axes.labelcolor":NAVY, "axes.edgecolor":"#CCD5DF",
        "grid.color":"#E7EDF3", "figure.facecolor":"white"})
    figs = {}
    def new(title, xlabel="Carga térmica acumulada (10³ kJ/h)", ylabel="Temperatura (°C)"):
        fig, ax = plt.subplots(figsize=(9, 5.4), layout="constrained")
        ax.set(title=title, xlabel=xlabel, ylabel=ylabel)
        ax.grid(alpha=.7)
        return fig, ax
    fig, ax = new("Curvas individuales · sentido real de cada corriente")
    for s in r["streams"]:
        x = np.array([s.duty, 0]) if s.hot else np.array([0, s.duty])
        y = [s.tin, s.tout]
        ax.plot(x/1000, y, "o-", color=HOT if s.hot else COLD, label=f"{s.id} · CP={s.cp:g}")
        ax.annotate("", xy=(x[1]/1000,y[1]), xytext=(x.mean()/1000,np.mean(y)),
                    arrowprops=dict(arrowstyle="->",color=HOT if s.hot else COLD))
    ax.legend(frameon=False); figs["01_individuales"] = fig
    for mode, title in [("originales","Compuestas sin ajuste · origen independiente"),
                        ("ajustadas","Compuestas ajustadas · recuperación máxima"),
                        ("desplazadas","Compuestas desplazadas · temperatura T*")]:
        fig, ax = new(title, ylabel="Temperatura desplazada T* (°C)" if mode=="desplazadas" else "Temperatura (°C)")
        for hot, color, label in [(True,HOT,"Caliente"),(False,COLD,"Fría")]:
            curve = r["hot"] if hot else r["cold"]
            xs = curve["h"].copy()
            if mode=="originales" and not hot: xs -= r["qc"]
            ys = curve["t"].copy()
            if mode=="desplazadas": ys += -r["dtmin"]/2 if hot else r["dtmin"]/2
            # Trazar cada intervalo activo evita inventar calor en huecos de T.
            first = True
            for k, cp in enumerate(curve["cp"]):
                if cp > 0 and xs[k+1] > xs[k]+1e-10:
                    ax.plot(xs[k:k+2]/1000,ys[k:k+2],"o-",color=color,
                            label=label if first else None, linewidth=2.5, markersize=3)
                    first = False
        if mode!="originales":
            ax.axvspan(0,r["qc"]/1000,color=COLD,alpha=.08)
            ax.axvspan(r["qhot"]/1000,(r["qc"]+r["qcold"])/1000,color=HOT,alpha=.08)
            ax.text(.02,.97,f"QHmin = {r['qh']:,.0f} kJ/h\nQCmin = {r['qc']:,.0f} kJ/h",
                    transform=ax.transAxes,va="top",bbox=dict(facecolor="white",edgecolor="none",alpha=.9))
            for p in r["pinches"]:
                if p["kind"] != "interior":
                    continue
                xp=float(np.interp(p["T_hot"],r["hot"]["t"],r["hot"]["h"]))/1000
                yh=p["T_star"] if mode=="desplazadas" else p["T_hot"]
                yc=p["T_star"] if mode=="desplazadas" else p["T_cold"]
                ax.plot([xp,xp],[yc,yh],color=NAVY,marker="o",ms=4,zorder=5)
                ax.annotate(f"Pinch {p['T_hot']:g} / {p['T_cold']:g} °C",
                            (xp,(yh+yc)/2),xytext=(12,15),textcoords="offset points",
                            fontsize=9,color=NAVY)
        ax.legend(frameon=False,loc="lower right")
        figs[{"originales":"02_compuestas_originales","ajustadas":"03_compuestas_ajustadas",
              "desplazadas":"04_compuestas_desplazadas"}[mode]] = fig
    fig, ax = new("Gran curva compuesta · cascada de calor",
                  "Calor residual (10³ kJ/h)","Temperatura desplazada T* (°C)")
    ax.plot(r["residual"]/1000,r["levels"],"o-",color=GREEN,linewidth=2.5)
    for p in r["pinches"]:
        ax.scatter([0],[p["T_star"]],color=NAVY,zorder=4)
        ax.annotate(f" T*={p['T_star']:g} °C",(0,p["T_star"]),xytext=(8,5),textcoords="offset points")
    figs["05_gran_compuesta"] = fig
    fig, ax = new("Temperaturas de suministro y objetivo", "Temperatura (°C)", "Corriente")
    for i,s in enumerate(r["streams"]):
        ax.annotate("", xy=(s.tout,i),xytext=(s.tin,i),
                    arrowprops=dict(arrowstyle="->",lw=3,color=HOT if s.hot else COLD))
        ax.text(s.tin,i+.16,f"{s.tin:g}°",ha="center",fontsize=9)
        ax.text(s.tout,i-.23,f"{s.tout:g}°",ha="center",fontsize=9)
    ax.set_yticks(range(len(r["streams"])),[s.id for s in r["streams"]])
    # Las flechas creadas con annotate no actualizan los límites de datos.
    tlo=min(s.low for s in r["streams"]); thi=max(s.high for s in r["streams"])
    pad=max(5,(thi-tlo)*.12)
    ax.set_xlim(tlo-pad,thi+pad); ax.set_ylim(-.6,len(r["streams"])-.4)
    figs["06_temperaturas"] = fig
    ds = np.linspace(0,max(40,2*r["dtmin"]),41)
    sens=sensitivity(r["streams"],ds)
    fig,ax=new("Sensibilidad al ΔT mínimo","ΔT mínimo (K)","Carga térmica (10³ kJ/h)")
    for key,label,color in [("QHmin","Calentamiento",HOT),("QCmin","Refrigeración",COLD),("Qrec","Recuperación",GREEN)]:
        ax.plot(sens.DTmin,sens[key]/1000,label=label,color=color,linewidth=2.3)
    ax.axvline(r["dtmin"],color=NAVY,ls="--",alpha=.5); ax.legend(frameon=False)
    figs["07_sensibilidad"] = fig
    # Mapa agregado de energía: no simula la topología ni oculta las ramas.
    hi=[i for i,s in enumerate(r["streams"]) if s.hot]
    ci=[i for i,s in enumerate(r["streams"]) if not s.hot]
    if hi and ci:
        fig,ax=plt.subplots(figsize=(9,5.4),layout="constrained")
        mat=r["matrix"][np.ix_(hi,ci)]/1000
        im=ax.imshow(mat,cmap="Blues",aspect="auto",vmin=0)
        for (i,j),v in np.ndenumerate(mat):
            ax.text(j,i,f"{v:,.1f}",ha="center",va="center",color="white" if v>mat.max()*.6 else NAVY)
        ax.set_xticks(range(len(ci)),[r["streams"][i].id for i in ci])
        ax.set_yticks(range(len(hi)),[r["streams"][i].id for i in hi])
        ax.set(title="Intercambios agregados entre corrientes",xlabel="Receptora fría",ylabel="Donante caliente")
        fig.colorbar(im,ax=ax,label="10³ kJ/h")
        figs["08_matriz_intercambios"]=fig
    # Red por tramos: cada columna es un intervalo entálpico. Una línea
    # vertical representa una pareja de ramas, no una corriente sin dividir.
    active=[b for b in r["blocks"] if b["q"]>r["tolerance"] and (b["hot"] or b["cold"])]
    fig,ax=plt.subplots(figsize=(max(10,len(active)*1.1),max(4,len(r["streams"])*.95)),layout="constrained")
    for i,s in enumerate(r["streams"]):
        color=HOT if s.hot else COLD
        ax.axhline(i,color=color,alpha=.22,lw=2)
    xmap={b["index"]:k for k,b in enumerate(active)}
    for e in r["exchangers"]:
        x=xmap[e["tramo"]]; i=next(i for i,s in enumerate(r["streams"]) if s.id==e["caliente"])
        j=next(i for i,s in enumerate(r["streams"]) if s.id==e["fria"])
        # Distribuir todas las parejas del tramo en posiciones distintas.
        group=[v for v in r["exchangers"] if v["tramo"]==e["tramo"]]
        local_index=next(k for k,v in enumerate(group) if v["equipo"]==e["equipo"])
        x += -.32 + .64*(local_index+1)/(len(group)+1)
        ax.plot([x,x],[i,j],"o-",color=GREEN,ms=5,alpha=.75)
        ax.text(x+.025,(i+j)/2,e["equipo"],fontsize=7,rotation=90,va="center")
    for urow in r["utilities"]:
        i=next(i for i,s in enumerate(r["streams"]) if s.id==urow["corriente"])
        x=xmap[urow["tramo"]]
        ax.scatter(x,i,marker="s",s=90,color=COLD if urow["servicio"]=="Refrigeración" else HOT,zorder=6)
    ax.set_yticks(range(len(r["streams"])),[s.id for s in r["streams"]])
    ax.set_xticks(range(len(active)),[f"{b['elo']/1000:g}–{b['ehi']/1000:g}" for b in active],rotation=35,ha="right")
    for p in r["pinches"]:
        if p["kind"]!="interior":continue
        hp=float(np.interp(p["T_hot"],r["hot"]["t"],r["hot"]["h"]))
        for k,b in enumerate(active):
            if abs(b["elo"]-hp)<r["tolerance"]:
                ax.axvline(k-.5,color=NAVY,ls="--",lw=1,alpha=.6)
                ax.text(k-.47,len(r["streams"])-.7,"Pinch",color=NAVY,fontsize=9)
                break
    ax.set(title="Red por tramos · parejas de ramas y servicios",
           xlabel="Intervalo de entalpía (10³ kJ/h); consultar fracciones en la tabla de equipos")
    ax.margins(.07,.2); figs["09_red_por_tramos"]=fig
    return figs


def export_all(r, output_dir):
    """Guarda Excel editable, tablas CSV, nueve gráficos y resumen JSON."""
    from excel_export import export_excel
    out=Path(output_dir); out.mkdir(parents=True,exist_ok=True)
    export_excel(r,out/"Pinch_Analisis.xlsx")
    for name,df in tables(r).items():
        df.to_csv(out/f"{name}.csv",index=name=="matriz_Q",encoding="utf-8-sig")
    for name,fig in figures(r).items():
        fig.savefig(out/f"{name}.png",dpi=160)
        fig.savefig(out/f"{name}.svg")
        plt.close(fig)
    summary={k:r[k] for k in ["dtmin","qh","qc","qhot","qcold","recovered","pinches","balance_errors"]}
    (out/"resumen.json").write_text(json.dumps(summary,indent=2,ensure_ascii=False),encoding="utf-8")
    return out


def main():
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--csv",type=Path,help="CSV de entrada; columnas id,cp,tin,tout")
    parser.add_argument("--excel",type=Path,help="Reimportar Entradas de un Excel de Pinch Lab")
    parser.add_argument("--dtmin",type=float,default=None,help="ΔT mínimo en K (por defecto: 10)")
    parser.add_argument("--u",type=float,default=None,help="U en kJ/(h m² K); 0 omite áreas")
    parser.add_argument("--out",type=Path,default=Path("resultados"))
    a=parser.parse_args()
    if a.csv and a.excel: parser.error("Elige --csv o --excel, no ambos.")
    streams,dt_file,u_file=read_streams(a.csv or a.excel) if a.csv or a.excel else (EXAMPLE,None,None)
    dt=a.dtmin if a.dtmin is not None else dt_file if dt_file is not None else 10.
    u=a.u if a.u is not None else u_file if u_file is not None else 0.
    try:
        r=analyse(streams,dt,u)
        out=export_all(r,a.out)
    except (ValueError,ArithmeticError) as exc:
        parser.exit(2,f"Datos no válidos: {exc}\n")
    print(f"QHmin: {r['qh']:,.2f} kJ/h = {r['qh']/3600:.3f} kW")
    print(f"QCmin: {r['qc']:,.2f} kJ/h = {r['qc']/3600:.3f} kW")
    print(f"Recuperación: {r['recovered']:,.2f} kJ/h")
    print(f"Pinches: {r['pinches']}")
    print(f"Archivos guardados en: {out.resolve()}")


if __name__=="__main__":
    main()
